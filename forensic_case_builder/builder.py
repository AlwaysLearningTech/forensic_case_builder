"""Builder module for generating Declaration document."""

import logging
import sqlite3
from pathlib import Path
from datetime import datetime
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

logger = logging.getLogger(__name__)


def build_case(db_path, output_dir):
    """Generate Declaration document in DOCX format."""
    conn = sqlite3.connect(db_path)
    cursor = conn.cursor()
    
    # Get summary statistics
    cursor.execute("SELECT COUNT(*) FROM evidence")
    total_files = cursor.fetchone()[0]
    
    cursor.execute("SELECT COUNT(*) FROM extracted_content")
    extracted_files = cursor.fetchone()[0]
    
    cursor.execute("SELECT SUM(file_size) FROM evidence")
    total_size = cursor.fetchone()[0] or 0
    
    cursor.execute("SELECT COUNT(DISTINCT file_type) FROM evidence")
    file_types = cursor.fetchone()[0]
    
    # Get evidence details
    cursor.execute("""
        SELECT 
            e.id,
            e.original_path,
            e.sha256,
            e.file_size,
            e.file_type,
            e.scan_timestamp
        FROM evidence e
        ORDER BY e.id
    """)
    
    evidence_list = cursor.fetchall()
    
    # Create Document
    doc = Document()
    
    # Title
    title = doc.add_heading('Declaration of Evidence Collection', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Metadata
    doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph()
    
    # Executive Summary
    doc.add_heading('Executive Summary', level=1)
    doc.add_paragraph(f"Total Evidence Files: {total_files}")
    doc.add_paragraph(f"Files with Extracted Content: {extracted_files}")
    doc.add_paragraph(f"Total Size: {total_size:,} bytes ({total_size / (1024*1024):.2f} MB)")
    doc.add_paragraph(f"Distinct File Types: {file_types}")
    doc.add_paragraph()
    
    # Declaration Statement
    doc.add_heading('Declaration', level=1)
    declaration_text = (
        "I hereby declare that the evidence listed in this document has been collected "
        "and processed in accordance with established forensic procedures. All original "
        "files remain unmodified and are preserved in their original location. Copies "
        "have been made to a secure processing area for analysis. Each file has been "
        "verified using SHA256 cryptographic hashing to ensure integrity."
    )
    doc.add_paragraph(declaration_text)
    doc.add_paragraph()
    
    # Evidence List
    doc.add_heading('Evidence Inventory', level=1)
    
    for evidence_id, orig_path, sha256, file_size, file_type, scan_time in evidence_list:
        # Add evidence entry
        p = doc.add_paragraph()
        p.add_run(f"Evidence ID: {evidence_id}").bold = True
        
        doc.add_paragraph(f"Original Path: {orig_path}", style='List Bullet')
        doc.add_paragraph(f"SHA256: {sha256}", style='List Bullet')
        doc.add_paragraph(f"File Size: {file_size:,} bytes", style='List Bullet')
        doc.add_paragraph(f"File Type: {file_type}", style='List Bullet')
        doc.add_paragraph(f"Scan Time: {scan_time}", style='List Bullet')
        doc.add_paragraph()
    
    # Signature Section
    doc.add_page_break()
    doc.add_heading('Certification', level=1)
    doc.add_paragraph()
    doc.add_paragraph("Forensic Analyst Signature: ___________________________")
    doc.add_paragraph()
    doc.add_paragraph("Date: ___________________________")
    doc.add_paragraph()
    doc.add_paragraph("Case Number: ___________________________")
    
    # Save document
    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    doc_path = output_path / f"declaration_{timestamp}.docx"
    doc.save(doc_path)
    
    conn.close()
    
    logger.info(f"Declaration saved to {doc_path} ({total_files} evidence items)")
