"""Reporter module for generating Evidence Log and timeline reports."""

import logging
import sqlite3
import csv
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

logger = logging.getLogger(__name__)


def generate_evidence_log_xlsx(db_path, output_path):
    """Generate Evidence Log in XLSX format."""
    conn = sqlite3.connect(db_path)
    cursor = conn.cursor()
    
    # Query evidence data
    cursor.execute("""
        SELECT 
            e.id,
            e.original_path,
            e.processing_path,
            e.sha256,
            e.file_size,
            e.file_type,
            e.created_time,
            e.modified_time,
            e.scan_timestamp
        FROM evidence e
        ORDER BY e.scan_timestamp
    """)
    
    rows = cursor.fetchall()
    
    # Create Excel workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Evidence Log"
    
    # Headers
    headers = [
        "ID", "Original Path", "Processing Path", "SHA256", 
        "File Size (bytes)", "File Type", "Created Time", 
        "Modified Time", "Scan Timestamp"
    ]
    
    # Style headers
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
    
    # Data rows
    for row_idx, row_data in enumerate(rows, 2):
        for col_idx, value in enumerate(row_data, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    
    # Adjust column widths
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column].width = adjusted_width
    
    wb.save(output_path)
    conn.close()
    
    logger.info(f"Evidence log saved to {output_path} ({len(rows)} entries)")


def generate_timeline_csv(db_path, output_path):
    """Generate timeline in CSV format."""
    conn = sqlite3.connect(db_path)
    cursor = conn.cursor()
    
    # Query timeline data
    cursor.execute("""
        SELECT 
            e.id,
            e.original_path,
            e.file_type,
            e.created_time,
            e.modified_time,
            e.accessed_time,
            e.scan_timestamp
        FROM evidence e
        ORDER BY e.modified_time
    """)
    
    rows = cursor.fetchall()
    
    # Write CSV
    with open(output_path, 'w', newline='', encoding='utf-8') as csvfile:
        writer = csv.writer(csvfile)
        writer.writerow([
            "Evidence ID", "File Path", "File Type", 
            "Created Time", "Modified Time", "Accessed Time", 
            "Scan Timestamp"
        ])
        writer.writerows(rows)
    
    conn.close()
    
    logger.info(f"Timeline CSV saved to {output_path} ({len(rows)} entries)")


def generate_timeline_xlsx(db_path, output_path):
    """Generate timeline in XLSX format."""
    conn = sqlite3.connect(db_path)
    cursor = conn.cursor()
    
    # Query timeline data
    cursor.execute("""
        SELECT 
            e.id,
            e.original_path,
            e.file_type,
            e.created_time,
            e.modified_time,
            e.accessed_time,
            e.scan_timestamp
        FROM evidence e
        ORDER BY e.modified_time
    """)
    
    rows = cursor.fetchall()
    
    # Create Excel workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Timeline"
    
    # Headers
    headers = [
        "Evidence ID", "File Path", "File Type", 
        "Created Time", "Modified Time", "Accessed Time", 
        "Scan Timestamp"
    ]
    
    # Style headers
    header_fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
    
    # Data rows
    for row_idx, row_data in enumerate(rows, 2):
        for col_idx, value in enumerate(row_data, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    
    # Adjust column widths
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column].width = adjusted_width
    
    wb.save(output_path)
    conn.close()
    
    logger.info(f"Timeline XLSX saved to {output_path} ({len(rows)} entries)")


def generate_reports(db_path, output_dir):
    """Generate all reports: Evidence Log (XLSX) and Timeline (CSV/XLSX)."""
    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    
    # Evidence Log
    evidence_log_path = output_path / f"evidence_log_{timestamp}.xlsx"
    generate_evidence_log_xlsx(db_path, evidence_log_path)
    
    # Timeline CSV
    timeline_csv_path = output_path / f"timeline_{timestamp}.csv"
    generate_timeline_csv(db_path, timeline_csv_path)
    
    # Timeline XLSX
    timeline_xlsx_path = output_path / f"timeline_{timestamp}.xlsx"
    generate_timeline_xlsx(db_path, timeline_xlsx_path)
    
    logger.info("All reports generated successfully")
